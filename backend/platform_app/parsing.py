from pathlib import Path
from datetime import timedelta
from zipfile import ZipFile
from django.db import transaction
from django.utils import timezone
from .models import Asset,Record

MAX_TEXT=1000000
def extract(asset):
    file=Path(asset.file.path);suffix=file.suffix.lower();status='ready';warning='';parts=[];size=0
    mode=asset.record.payload.get('parseMode','auto');language=asset.record.payload.get('parseLanguage','zh')
    if mode in ['ocr','transcribe'] or suffix in ['.png','.jpg','.jpeg','.tif','.tiff','.bmp','.webp','.wav','.mp3','.mp4','.m4a','.webm','.mkv','.ogg','.flac']:
        from .media import process
        selected=mode if mode!='auto' else 'ocr' if suffix in ['.png','.jpg','.jpeg','.tif','.tiff','.bmp','.webp'] else 'transcribe'
        text,state,note,metadata=process(asset,selected,language);asset._processing_metadata=metadata;return text,state,note
    def add(value):
        nonlocal size
        if size>=MAX_TEXT:return False
        value=str(value);parts.append(value[:MAX_TEXT-size]);size+=len(value)
        return size<MAX_TEXT
    if suffix in ['.docx','.xlsx','.pptx']:
        with ZipFile(file) as archive:
            items=archive.infolist()
            if len(items)>20000 or sum(x.file_size for x in items)>100*1024*1024:
                return '', 'saved_only','原件已保存；解压后内容过大，等待人工处理。'
    if suffix in ['.txt','.md','.markdown','.tex','.bib','.csv','.tsv','.json','.py','.r','.m','.js','.srt','.vtt']:
        if file.stat().st_size>5*1024*1024:return '', 'saved_only','原件已保存；文本超过5MB，等待分块解析。'
        raw=file.read_bytes()
        try:add(raw.decode('utf-8-sig'))
        except UnicodeDecodeError:add(raw.decode('gb18030',errors='replace'))
    elif suffix=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(file)
        if reader.is_encrypted:return '', 'saved_only','加密PDF已保存，请提供可读取版本。'
        for page in reader.pages[:500]:
            if not add((page.extract_text() or '')+'\n\n'):break
        if not ''.join(parts).strip():
            from .media import process
            text,state,note,metadata=process(asset,'ocr',language);asset._processing_metadata=metadata;return text,state,note
        if len(reader.pages)>500:status='partial';warning='仅提取前500页。'
    elif suffix=='.docx':
        from docx import Document
        doc=Document(file)
        for p in doc.paragraphs:
            if not add(p.text+'\n'):break
        if size<MAX_TEXT:
            for table in doc.tables:
                for row in table.rows:
                    if not add(' | '.join(c.text for c in row.cells)+'\n'):break
                if size>=MAX_TEXT:break
    elif suffix=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(file,read_only=True,data_only=True)
        try:
            for sheet in book:
                add(sheet.title+'\n')
                for row in sheet.iter_rows(max_row=min(sheet.max_row or 10000,10000),max_col=min(sheet.max_column or 200,200),values_only=True):
                    if not add(' | '.join('' if c is None else str(c) for c in row)+'\n'):break
                if (sheet.max_row and sheet.max_row>10000) or (sheet.max_column and sheet.max_column>200):status='partial';warning='表格仅提取每页前10000行、200列。'
                if size>=MAX_TEXT:break
        finally:book.close()
    elif suffix=='.pptx':
        from pptx import Presentation
        for slide in Presentation(file).slides:
            if not add('\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)+'\n\n'):break
    else:return '', 'saved_only','原件已保存；此格式暂不自动提取正文。'
    if size>=MAX_TEXT:status='partial';warning='正文超过100万字符，已保留原件并截断检索文本。'
    return ''.join(parts),status,warning

def parse_asset(asset):
    try:text,status,warning=extract(asset)
    except Exception:text,status,warning='','failed','解析失败，原件仍保留，可在管理端将解析状态改为 pending 重试。'
    with transaction.atomic():
        current=Asset.objects.select_for_update().get(pk=asset.pk)
        if current.parse_status!='processing' or current.claimed_at!=asset.claimed_at:return
        row=Record.objects.select_for_update().get(pk=asset.record_id)
        row.payload.update(text=text,parseStatus=status,parseWarning=warning,**getattr(asset,'_processing_metadata',{}))
        row.revision+=1;row.save(update_fields=['payload','revision','updated_at'])
        from .retrieval import invalidate_resource
        invalidate_resource(row)
        current.parse_status=status;current.parse_error=warning;current.claimed_at=None;current.save()

def parse_pending(limit=5):
    Asset.objects.filter(parse_status='processing',claimed_at__lt=timezone.now()-timedelta(minutes=15)).update(parse_status='failed',parse_error='解析任务中断或超时，原件保留，请人工重试。')
    for key in list(Asset.objects.filter(parse_status='pending').values_list('pk',flat=True)[:limit]):
        with transaction.atomic():
            asset=Asset.objects.select_for_update().get(pk=key)
            if asset.parse_status!='pending':continue
            asset.parse_status='processing';asset.claimed_at=timezone.now();asset.save()
        parse_asset(asset)
